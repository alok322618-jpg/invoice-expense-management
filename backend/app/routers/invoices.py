import csv
import datetime
import io
import os
import uuid
from typing import List, Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session

from .. import auth as authlib
from .. import models, schemas
from ..database import get_db
from .vendors import resolve_vendor

router = APIRouter(prefix="/invoices", tags=["invoices"])


def _out(inv: models.Invoice) -> dict:
    return {
        "id": inv.id,
        "invoice_no": inv.invoice_no,
        "vendor": inv.vendor,
        "vendor_id": inv.vendor_id,
        "amount": inv.amount,
        "invoice_date": inv.invoice_date,
        "due_date": inv.due_date,
        "status": inv.status,
        "assignee_id": inv.assignee_id,
        "assignee_name": inv.assignee.name if inv.assignee else None,
        "account_code": inv.account_code or "",
        "coded_by_name": inv.coded_by.name if inv.coded_by else None,
        "coded_at": inv.coded_at,
        "created_by_id": inv.created_by_id,
        "created_by_name": inv.created_by.name if inv.created_by else None,
        "attachment": inv.attachment or "",
        "event_id": inv.event_id,
        "event_name": inv.event.name if inv.event else None,
        "remarks": inv.remarks or "",
        "history": inv.history or [],
        "created_at": inv.created_at,
    }


def _log(inv: models.Invoice, by: str, action: str, remarks: str = ""):
    hist = list(inv.history or [])
    hist.append({
        "at": datetime.datetime.utcnow().isoformat(),
        "by": by,
        "action": action,
        "remarks": remarks,
    })
    inv.history = hist


@router.get("", response_model=List[schemas.InvoiceOut])
def list_invoices(
    status: Optional[str] = None,
    search: Optional[str] = None,
    event_id: Optional[int] = None,
    db: Session = Depends(get_db),
    user: models.User = Depends(authlib.get_current_user),
):
    q = db.query(models.Invoice)
    if user.role == "employee":
        q = q.filter_by(created_by_id=user.id)
    if status:
        q = q.filter_by(status=status)
    if event_id is not None:
        q = q.filter_by(event_id=event_id)
    if search:
        like = f"%{search}%"
        q = q.filter((models.Invoice.invoice_no.like(like)) | (models.Invoice.vendor.like(like)))
    rows = q.order_by(models.Invoice.id.desc()).all()
    return [_out(r) for r in rows]


@router.post("", response_model=schemas.InvoiceOut)
def create_invoice(
    body: schemas.InvoiceCreateIn,
    db: Session = Depends(get_db),
    user: models.User = Depends(authlib.require_roles("admin", "manager", "finance", "employee")),
):
    if db.query(models.Invoice).filter_by(invoice_no=body.invoice_no).first():
        raise HTTPException(status.HTTP_409_CONFLICT, "Invoice number already exists")
    vendor = resolve_vendor(db, body.vendor_id, body.vendor)
    event = None
    if body.event_id is not None:
        event = db.query(models.Event).filter_by(id=body.event_id).first()
        if not event:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Event not found")
    inv = models.Invoice(
        invoice_no=body.invoice_no,
        vendor=vendor.name,
        vendor_id=vendor.id,
        amount=body.amount,
        invoice_date=body.invoice_date,
        due_date=body.due_date,
        status="captured",
        created_by_id=user.id,
        event_id=event.id if event else None,
    )
    _log(inv, user.username, "captured")
    db.add(inv)
    db.commit()
    db.refresh(inv)
    return _out(inv)


def _parse_date(value):
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _read_rows(filename, content):
    name = (filename or "").lower()
    if name.endswith(".csv"):
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        rows = []
        for row in reader:
            rows.append({
                (k or "").strip().lower(): (v.strip() if isinstance(v, str) else v)
                for k, v in row.items() if k
            })
        return rows
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        data = list(ws.iter_rows(values_only=True))
        if not data:
            return []
        headers = [str(h).strip().lower() if h is not None else "" for h in data[0]]
        rows = []
        for r in data[1:]:
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in r):
                continue
            rows.append({headers[i]: r[i] for i in range(min(len(headers), len(r)))})
        return rows
    raise HTTPException(status.HTTP_400_BAD_REQUEST, "Only .csv and .xlsx files are supported")


@router.post("/bulk-upload", response_model=schemas.BulkUploadOut)
def bulk_upload(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: models.User = Depends(authlib.require_roles("admin", "manager", "finance")),
):
    content = file.file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "File too large (max 5 MB)")
    rows = _read_rows(file.filename, content)
    if len(rows) > 1000:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Too many rows (max 1000 per upload)")
    existing = {n for (n,) in db.query(models.Invoice.invoice_no).all()}
    seen = set()
    added = 0
    skipped = []
    for idx, row in enumerate(rows, start=2):
        inv_no = str(row.get("invoice_no") or "").strip()
        vendor = str(row.get("vendor") or "").strip()
        inv_date = _parse_date(row.get("invoice_date"))
        due_date = _parse_date(row.get("due_date"))
        amount = 0.0
        try:
            amount = float(row.get("amount"))
        except (TypeError, ValueError):
            amount = 0.0
        reason = ""
        if not inv_no:
            reason = "missing invoice_no"
        elif not vendor:
            reason = "missing vendor"
        elif inv_no in existing or inv_no in seen:
            reason = "duplicate invoice_no"
        elif amount <= 0:
            reason = "invalid amount"
        elif not inv_date:
            reason = "invalid invoice_date"
        if reason:
            skipped.append({"row": idx, "invoice_no": inv_no, "reason": reason})
            continue
        vendor = resolve_vendor(db, None, vendor)
        inv = models.Invoice(
            invoice_no=inv_no,
            vendor=vendor.name,
            vendor_id=vendor.id,
            amount=amount,
            invoice_date=inv_date,
            due_date=due_date,
            status="captured",
            created_by_id=user.id,
        )
        _log(inv, user.username, "bulk uploaded")
        db.add(inv)
        seen.add(inv_no)
        existing.add(inv_no)
        added += 1
        if added % 200 == 0:
            db.commit()
    db.commit()
    return {"added": added, "skipped": skipped, "total": len(rows)}


@router.get("/{invoice_id}", response_model=schemas.InvoiceOut)
def get_invoice(
    invoice_id: int,
    db: Session = Depends(get_db),
    user: models.User = Depends(authlib.get_current_user),
):
    inv = db.query(models.Invoice).filter_by(id=invoice_id).first()
    if not inv:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Invoice not found")
    if user.role == "employee" and inv.created_by_id != user.id:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Not allowed for your role")
    return _out(inv)


@router.patch("/{invoice_id}/assign", response_model=schemas.InvoiceOut)
def assign_invoice(
    invoice_id: int,
    body: schemas.AssignIn,
    db: Session = Depends(get_db),
    user: models.User = Depends(authlib.require_roles("admin", "manager")),
):
    inv = db.query(models.Invoice).filter_by(id=invoice_id).first()
    if not inv:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Invoice not found")
    assignee = db.query(models.User).filter_by(id=body.assignee_id, status="active").first()
    if not assignee:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Assignee not found or inactive")
    inv.assignee_id = assignee.id
    inv.status = "assigned"
    _log(inv, user.username, f"assigned to {assignee.name}")
    db.commit()
    db.refresh(inv)
    return _out(inv)


@router.patch("/{invoice_id}/verify", response_model=schemas.InvoiceOut)
def verify_invoice(
    invoice_id: int,
    body: schemas.VerifyIn,
    db: Session = Depends(get_db),
    user: models.User = Depends(authlib.require_roles("admin", "manager", "finance")),
):
    if body.decision not in ("approve", "reject"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "decision must be approve or reject")
    if body.decision == "reject" and not body.remarks.strip():
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Rejection needs a reason")
    inv = db.query(models.Invoice).filter_by(id=invoice_id).first()
    if not inv:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Invoice not found")
    if inv.status not in ("assigned", "in_verification", "captured"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Cannot verify invoice in status '{inv.status}'")
    if body.decision == "approve" and not (inv.account_code or "").strip():
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Invoice must be coded before approval")
    inv.status = "approved" if body.decision == "approve" else "rejected"
    inv.remarks = body.remarks
    _log(inv, user.username, inv.status, body.remarks)
    db.commit()
    db.refresh(inv)
    return _out(inv)


@router.patch("/{invoice_id}/code", response_model=schemas.InvoiceOut)
def code_invoice(
    invoice_id: int,
    body: schemas.CodeIn,
    db: Session = Depends(get_db),
    user: models.User = Depends(authlib.require_roles("admin", "manager", "finance")),
):
    code = (body.account_code or "").strip()
    if not code:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Account code is required")
    inv = db.query(models.Invoice).filter_by(id=invoice_id).first()
    if not inv:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Invoice not found")
    if inv.status not in ("captured", "assigned", "in_verification"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Cannot code invoice in status '{inv.status}'")
    ac = db.query(models.AccountCode).filter(
        models.AccountCode.code.ilike(code),
        models.AccountCode.status == "active",
    ).first()
    if not ac:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Unknown or inactive account code")
    inv.account_code = ac.code
    inv.coded_by_id = user.id
    inv.coded_at = datetime.datetime.utcnow()
    _log(inv, user.username, f"coded as {ac.code}", body.remarks)
    db.commit()
    db.refresh(inv)
    return _out(inv)


@router.patch("/{invoice_id}/event", response_model=schemas.InvoiceOut)
def link_event(
    invoice_id: int,
    body: schemas.EventLinkIn,
    db: Session = Depends(get_db),
    user: models.User = Depends(authlib.require_roles("admin", "manager", "finance")),
):
    inv = db.query(models.Invoice).filter_by(id=invoice_id).first()
    if not inv:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Invoice not found")
    if body.event_id is not None:
        event = db.query(models.Event).filter_by(id=body.event_id).first()
        if not event:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Event not found")
        inv.event_id = event.id
        _log(inv, user.username, f"linked to event {event.name}")
    else:
        inv.event_id = None
        _log(inv, user.username, "unlinked from event")
    db.commit()
    db.refresh(inv)
    return _out(inv)


@router.patch("/{invoice_id}/send-back", response_model=schemas.InvoiceOut)
def send_back_invoice(
    invoice_id: int,
    db: Session = Depends(get_db),
    user: models.User = Depends(authlib.require_roles("admin", "manager")),
):
    inv = db.query(models.Invoice).filter_by(id=invoice_id, status="rejected").first()
    if not inv:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Rejected invoice not found")
    inv.status = "assigned"
    _log(inv, user.username, "sent back for correction")
    db.commit()
    db.refresh(inv)
    return _out(inv)


UPLOAD_DIR = os.environ.get("UPLOAD_DIR", "./uploads")
ALLOWED_EXT = {".png", ".jpg", ".jpeg", ".pdf", ".webp"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024


def _upload_dir():
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    return UPLOAD_DIR


def _get_for_upload(invoice_id: int, db: Session, user: models.User) -> models.Invoice:
    inv = db.query(models.Invoice).filter_by(id=invoice_id).first()
    if not inv:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Invoice not found")
    if user.role == "employee" and inv.created_by_id != user.id:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Not allowed for your role")
    return inv


@router.post("/{invoice_id}/attachment", response_model=schemas.InvoiceOut)
def upload_attachment(
    invoice_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: models.User = Depends(authlib.require_roles("admin", "manager", "finance", "employee")),
):
    inv = _get_for_upload(invoice_id, db, user)
    ext = os.path.splitext((file.filename or "").strip())[1].lower()
    if ext not in ALLOWED_EXT:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Only PNG, JPG, PDF or WEBP files are allowed")
    content = file.file.read()
    if not content:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Empty file")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "File too large (max 10 MB)")
    safe = f"{invoice_id}_{uuid.uuid4().hex}{ext}"
    with open(os.path.join(_upload_dir(), safe), "wb") as f:
        f.write(content)
    if inv.attachment:
        old = os.path.join(_upload_dir(), os.path.basename(inv.attachment))
        if os.path.isfile(old):
            os.remove(old)
    inv.attachment = safe
    _log(inv, user.username, "attachment uploaded")
    db.commit()
    db.refresh(inv)
    return _out(inv)


@router.get("/{invoice_id}/attachment")
def download_attachment(
    invoice_id: int,
    db: Session = Depends(get_db),
    user: models.User = Depends(authlib.get_current_user),
):
    inv = _get_for_upload(invoice_id, db, user)
    if not inv.attachment:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "No attachment for this invoice")
    path = os.path.join(_upload_dir(), os.path.basename(inv.attachment))
    if not os.path.isfile(path):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Attachment file missing")
    return FileResponse(path)
